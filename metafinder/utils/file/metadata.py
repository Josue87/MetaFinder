from docx import Document
import pikepdf
from openpyxl import load_workbook
from pptx import Presentation
from datetime import datetime


def _get_properties(prop):
    metadata = {}
    metadata["Author"] = prop.author
    metadata["Comments"] = prop.comments
    metadata["Created"] = str(prop.created)
    metadata["Identifier"] = prop.identifier
    metadata["Keywords"] = prop.keywords
    metadata["Modified"] = str(prop.modified)
    metadata["Subject"] = prop.subject
    metadata["Title"] = prop.title

def extract_doc(document):
    doc = Document(document)
    prop = doc.core_properties
    return _get_properties(prop)

def extract_pdf(document):
    metadata = {}
    with pikepdf.open(document) as pdf:
        info = pdf.docinfo
        for meta in info:
            data = str(info.get(meta, ""))
            if data and data.startswith("D:"):
                try:
                    d = data.split("D:")[1].split("+")[0]
                    data = str(datetime.strptime(d,"%Y%m%d%H%M%S"))
                except:
                    pass
            metadata[meta[1:]] = data # [1:] avoid first element '/'
    return metadata


def extract_xls(document):
    wb = load_workbook(document)
    prop = wb.properties
    return _get_properties(prop)

def extract_ppt(document):
    pptx_presentation = Presentation(document)
    prop = pptx_presentation.core_properties
    return _get_properties(prop)


def remove_indirect_object(metadata):
    new_metadata = {}
    for m in metadata:
        value =  metadata[m]
        if "IndirectObject(" not in str(value) and str(value) != "":
            new_metadata[m] = value
    return new_metadata


def extract_metadata(document):
    try:
        if document.endswith("pdf"):
            metadata = extract_pdf(document)
        elif document.endswith("doc") or document.endswith("docx"):
            metadata = extract_doc(document)
        elif document.endswith("xls") or document.endswith("xlsx"):
            metadata = extract_xls(document)
        elif document.endswith("ppt") or document.endswith("pptx"):
            metadata = extract_ppt(document)
        else: 
            metadata = {}
    except:
        metadata = {}
  
    return remove_indirect_object(metadata)




